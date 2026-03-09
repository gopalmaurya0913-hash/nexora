from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()

# Slide 1: Title (The Slogan)
slide_layout = prs.slide_layouts[0] # Title slide
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Nexoraa: Turning social media chaos into digital lifelines."
subtitle.text = "AI-Powered Disaster Relief Coordinator\nPresented by: Gopal Maurya"

# Slide 2: The Problem
slide_layout = prs.slide_layouts[1] # Title and Content
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
# Adjust body width to make room for image on the right
body.width = Inches(5)

title.text = "1. The Chaos of a Disaster"
tf = body.text_frame
tf.text = "When disaster strikes, information is chaotic and overwhelming."

p = tf.add_paragraph()
p.text = "Problem: Social media is flooded with unstructured cries for help (e.g., 'Trapped on roof!', 'Need water!')."
p.level = 1

p = tf.add_paragraph()
p.text = "The Bottleneck: Rescue teams receive too much noisy data to process manually."
p.level = 1

p = tf.add_paragraph()
p.text = "The Cost: Delayed response times cost lives because the noise hides the actual urgent needs."
p.level = 1

# Add disaster image to the right
img_path_1 = r"C:\Users\gopal\.antigravity\anti\nexoraa\disaster_real.jpg"
slide.shapes.add_picture(img_path_1, Inches(5.5), Inches(2), width=Inches(4))

# Slide 3: The Solution (Nexoraa)
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
body.width = Inches(5)

title.text = "2. The Solution: Nexoraa"
tf = body.text_frame
tf.text = "An AI command center that organizes the chaos instantly."

p = tf.add_paragraph()
p.text = "AI Categorization: Natural Language Processing reads feeds and tags messages by urgency (Medical, Rescue, Food)."
p.level = 1

p = tf.add_paragraph()
p.text = "Dynamic Geocoding: Extracts locations (like 'Surat') and uses Geopy to map them instantly to precise real-world GPS coordinates."
p.level = 1

p = tf.add_paragraph()
p.text = "Live Dashboard: Rescue teams see a real-time, dark-mode visual heatmap to prioritize field deployments immediately."
p.level = 1

# Add Nexoraa app screenshot to the right
img_path_2 = r"C:\Users\gopal\.gemini\antigravity\brain\6a9f0743-1c3c-455a-b4ff-67f3fe73657e\.system_generated\click_feedback\click_feedback_1772879799918.png"
slide.shapes.add_picture(img_path_2, Inches(5.5), Inches(2), width=Inches(4))

# Slide 4: The Impact
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.shapes.placeholders[1]
body.width = Inches(5)

title.text = "3. The Future & Impact"
tf = body.text_frame
tf.text = "Seconds save lives."

p = tf.add_paragraph()
p.text = "Speed: Automating the triage of SOS messages reduces decision time from hours to milliseconds."
p.level = 1

p = tf.add_paragraph()
p.text = "Synchronization: Cloud-deployed (Render) command center ensures teams in the field and at HQ are perfectly synced."
p.level = 1

p = tf.add_paragraph()
p.text = "Result: Faster deployments, smarter operations, and more people brought home safely."
p.level = 1

# Add rescue helicopter image to the right
img_path_3 = r"C:\Users\gopal\.antigravity\anti\nexoraa\rescue_real.jpg"
slide.shapes.add_picture(img_path_3, Inches(5.5), Inches(2), width=Inches(4))
# Save the presentation
prs.save("Nexoraa_Presentation_V2.pptx")
print("Presentation successfully created as V2!")
