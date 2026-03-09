import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()

def apply_dark_theme(slide):
    # Set background color to dark blue/grey
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(15, 23, 42) # slate-900 like dark mode

def style_text(shape, color=RGBColor(255, 255, 255), is_title=False):
    if not shape.has_text_frame:
        return
    for p in shape.text_frame.paragraphs:
        for run in p.runs:
            run.font.color.rgb = color
            if is_title:
                run.font.bold = True
                run.font.size = Pt(40)
                run.font.color.rgb = RGBColor(0, 255, 255) # Cyan highlight for titles
            else:
                run.font.size = Pt(24)

# Slide 1: Title
slide_layout = prs.slide_layouts[6] # Blank
slide = prs.slides.add_slide(slide_layout)
apply_dark_theme(slide)

txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(2))
tf = txBox.text_frame
p = tf.add_paragraph()
p.text = "Nexoraa"
p.alignment = PP_ALIGN.CENTER
p.runs[0].font.bold = True
p.runs[0].font.size = Pt(60)
p.runs[0].font.color.rgb = RGBColor(0, 255, 255)

p2 = tf.add_paragraph()
p2.text = "Turning Social Media Chaos into Digital Lifelines"
p2.alignment = PP_ALIGN.CENTER
p2.runs[0].font.size = Pt(30)
p2.runs[0].font.color.rgb = RGBColor(255, 255, 255)

p3 = tf.add_paragraph()
p3.text = "\nBy: Gopal Maurya"
p3.alignment = PP_ALIGN.CENTER
p3.runs[0].font.size = Pt(20)
p3.runs[0].font.color.rgb = RGBColor(200, 200, 200)

# Slide 2: The Problem
slide_layout = prs.slide_layouts[5] # Title only
slide = prs.slides.add_slide(slide_layout)
apply_dark_theme(slide)
title = slide.shapes.title
title.text = "1. When Disaster Strikes"
style_text(title, is_title=True)

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4.5), Inches(4))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.add_paragraph()
p.text = "• People use phones to cry for help.\n\n• The internet floods with thousands of messy SOS messages.\n\n• Humans CANNOT read them fast enough.\n\n• Because of this noise, rescue teams don't know where to go first."
style_text(txBox)

img_path_1 = r"C:\Users\gopal\.antigravity\anti\nexoraa\disaster_real.jpg"
if os.path.exists(img_path_1):
    slide.shapes.add_picture(img_path_1, Inches(5.2), Inches(1.8), width=Inches(4.5))

# Slide 3: The Big Idea
slide_layout = prs.slide_layouts[5] 
slide = prs.slides.add_slide(slide_layout)
apply_dark_theme(slide)
title = slide.shapes.title
title.text = "2. How AI Can Help"
style_text(title, is_title=True)

txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.add_paragraph()
p.text = "Instead of humans reading every tweet, what if a computer did it instantly?"
p.alignment = PP_ALIGN.CENTER
p.runs[0].font.bold = True

p2 = tf.add_paragraph()
p2.text = "\n1. AI Reads the text: 'Need water in Surat!'\n2. AI Understands it: Tags it as a 'High Urgency Water Need'.\n3. AI Finds the Location: Pinpoints 'Surat' on a real map automatically."
p2.alignment = PP_ALIGN.CENTER
style_text(txBox)

# Slide 4: The Solution (Nexoraa Setup)
slide_layout = prs.slide_layouts[5] 
slide = prs.slides.add_slide(slide_layout)
apply_dark_theme(slide)
title = slide.shapes.title
title.text = "3. Meet Nexoraa: The Dashboard"
style_text(title, is_title=True)

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4.5), Inches(4))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.add_paragraph()
p.text = "• A clean, dark-mode command center.\n\n• Live SOS Map: Drops glowing red pins exactly where help is needed.\n\n• Live Feed: Shows a simple list of emergencies in real-time.\n\n• Global Reach: Works anywhere in the world."
style_text(txBox)

img_path_2 = r"C:\Users\gopal\.gemini\antigravity\brain\6a9f0743-1c3c-455a-b4ff-67f3fe73657e\.system_generated\click_feedback\click_feedback_1772879799918.png"
if os.path.exists(img_path_2):
    slide.shapes.add_picture(img_path_2, Inches(5.2), Inches(2), width=Inches(4.5))

# Slide 5: The Impact
slide_layout = prs.slide_layouts[5] 
slide = prs.slides.add_slide(slide_layout)
apply_dark_theme(slide)
title = slide.shapes.title
title.text = "4. Why It Matters"
style_text(title, is_title=True)

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4.5), Inches(4))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.add_paragraph()
p.text = "• Seconds Save Lives: It instantly tells helicopters and ambulances where to go.\n\n• No messages get lost: Every cry for help is mapped.\n\n• Future of Rescue: Smarter, faster, and completely automated triage."
style_text(txBox)

img_path_3 = r"C:\Users\gopal\.antigravity\anti\nexoraa\rescue_real.jpg"
if os.path.exists(img_path_3):
    slide.shapes.add_picture(img_path_3, Inches(5.2), Inches(1.8), width=Inches(4.5))

# Save the presentation
prs.save("Nexoraa_Hackathon_Pitch.pptx")
print("Attractive 5-Slide Presentation successfully created!")
