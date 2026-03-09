import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
# Set slide dimensions to 16:9 widescreen
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

bg_image_path = r"C:\Users\gopal\.gemini\antigravity\brain\6a9f0743-1c3c-455a-b4ff-67f3fe73657e\tech_dark_background_1772987763122.png"

def add_background(slide):
    if os.path.exists(bg_image_path):
        slide.shapes.add_picture(bg_image_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
    else:
        # Fallback solid color
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(15, 23, 42)

def style_run(run, size_pt, bold=False, color=RGBColor(255, 255, 255), is_title=False):
    run.font.name = "Segoe UI"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    if is_title:
        run.font.color.rgb = RGBColor(0, 255, 255) # Cyan

def add_title(slide, text):
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(1))
    p = txBox.text_frame.add_paragraph()
    p.text = text
    style_run(p.runs[0], size_pt=45, bold=True, is_title=True)
    return txBox

def add_body(slide, text, left, top, width, height, font_size=28):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    paragraphs = text.split('\n')
    for i, line in enumerate(paragraphs):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = line
        if line.strip():
            style_run(p.runs[0], size_pt=font_size)
    return txBox

# --- Slide 1: Title ---
slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
add_background(slide)

txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(2))
tf = txBox.text_frame
p = tf.add_paragraph()
p.text = "Nexoraa"
p.alignment = PP_ALIGN.CENTER
style_run(p.runs[0], size_pt=90, bold=True, is_title=True)

p2 = tf.add_paragraph()
p2.text = "Turning Social Media Chaos into Digital Lifelines"
p2.alignment = PP_ALIGN.CENTER
style_run(p2.runs[0], size_pt=35)

p3 = tf.add_paragraph()
p3.text = "\nBy: Gopal Maurya"
p3.alignment = PP_ALIGN.CENTER
style_run(p3.runs[0], size_pt=24, color=RGBColor(200, 200, 200))


# --- Slide 2: The Problem ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_title(slide, "1. When Disaster Strikes")

body_text = (
    "• People use phones to cry for help.\n\n"
    "• The internet floods with thousands of messy SOS messages.\n\n"
    "• Humans CANNOT read them fast enough.\n\n"
    "• Because of this noise, rescue teams don't know where to go first."
)
add_body(slide, body_text, 0.5, 1.8, 6.5, 5, font_size=30)

img_path_1 = r"C:\Users\gopal\.antigravity\anti\nexoraa\disaster_real.jpg"
if os.path.exists(img_path_1):
    slide.shapes.add_picture(img_path_1, Inches(7.5), Inches(2), width=Inches(5.5))


# --- Slide 3: The Big Idea ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_title(slide, "2. How AI Can Help")

intro = "Instead of humans reading every tweet, what if a computer did it instantly?"
txIntro = add_body(slide, intro, 1, 2, 11.333, 1, font_size=35)
txIntro.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
txIntro.text_frame.paragraphs[0].runs[0].font.bold = True

steps_text = (
    "1. AI Reads the text: 'Need water in Surat!'\n\n"
    "2. AI Understands it: Tags it as a 'High Urgency Water Need'.\n\n"
    "3. AI Finds the Location: Pinpoints 'Surat' on a real map automatically."
)
add_body(slide, steps_text, 2, 3.5, 9.333, 3, font_size=32)


# --- Slide 4: The Solution (Nexoraa Setup) ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_title(slide, "3. Meet Nexoraa: The Dashboard")

dash_text = (
    "• A clean, dark-mode command center.\n\n"
    "• Live SOS Map: Drops glowing red pins exactly where help is needed.\n\n"
    "• Global Reach: Works anywhere in the world."
)
add_body(slide, dash_text, 0.5, 1.8, 6, 5, font_size=30)

# App image 1 (Main dashboard / History)
img_path_dash1 = r"C:\Users\gopal\.gemini\antigravity\brain\6a9f0743-1c3c-455a-b4ff-67f3fe73657e\.system_generated\click_feedback\click_feedback_1772879799918.png"
if os.path.exists(img_path_dash1):
    slide.shapes.add_picture(img_path_dash1, Inches(6.8), Inches(1), width=Inches(6.2))

# App image 2 (Zoomed out global map)
img_path_dash2 = r"C:\Users\gopal\.gemini\antigravity\brain\6a9f0743-1c3c-455a-b4ff-67f3fe73657e\.system_generated\click_feedback\click_feedback_1772879717284.png"
if os.path.exists(img_path_dash2):
    slide.shapes.add_picture(img_path_dash2, Inches(6.8), Inches(4.3), width=Inches(6.2))


# --- Slide 5: The Impact ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_title(slide, "4. Why It Matters")

impact_text = (
    "• Seconds Save Lives: It instantly tells helicopters and ambulances where to go.\n\n"
    "• No messages get lost: Every cry for help is mapped.\n\n"
    "• Future of Rescue: Smarter, faster, and completely automated triage."
)
add_body(slide, impact_text, 0.5, 1.8, 6.5, 5, font_size=30)

img_path_3 = r"C:\Users\gopal\.antigravity\anti\nexoraa\rescue_real.jpg"
if os.path.exists(img_path_3):
    slide.shapes.add_picture(img_path_3, Inches(7.5), Inches(2), width=Inches(5.5))

# Save the presentation
prs.save("Nexoraa_Hackathon_Pitch_V4.pptx")
print("Gorgeous 16:9 Presentation with background and multiple images created!")
